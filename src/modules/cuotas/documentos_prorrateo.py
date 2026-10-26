"""PDF y DOCX privados a partir del mismo reporte autorizado, sin consultar saldos."""

from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape


def tablas_reporte(reporte):
    tablas = []
    for edificio in reporte["edificios"]:
        filas = [["Código", "Concepto", "OBS", "Condominio", "Bruto edificio", "Cubierto moras", "Distribuible", "Unitario"]]
        total = Decimal("0.00")
        ajustes = [["Concepto", "Unidad", "Ajuste"]]
        numeros = {d["id"]: d["numero"] for d in edificio.get("departamentos", [])}
        for seccion in ("AACC", "EDIFICIO", "CRONOGRAMA_AACC"):
            subtotal = Decimal("0.00")
            filas.append(["", seccion.replace("_", " "), "", "", "", "", "", ""])
            for item in reporte["items"]:
                if item["seccion"] != seccion:
                    continue
                for a in item["asignaciones"]:
                    if a["edificio_id"] != edificio["id"]:
                        continue
                    condominio = "—" if item["metodo_asignacion"] == "DIRECTO_EDIFICIO" else (
                        ("-" if item["tipo_movimiento"] == "ABONO" else "") + item["monto_periodo"]
                    )
                    filas.append([item["codigo"], item["descripcion"], item.get("observacion_origen", ""), condominio,
                                  a.get("monto_bruto", a["monto"]), a.get("monto_cubierto_moras", "0.00"), a["monto"], a["unitario"]])
                    subtotal += Decimal(a["monto"])
                    for c in a.get("cuotas", []):
                        if Decimal(c["ajuste"]) != 0:
                            ajustes.append([item["codigo"], numeros.get(c["departamento_id"], c["departamento_id"]), c["ajuste"]])
            filas.append(["", "SUBTOTAL", "", "", "", "", f"{subtotal:.2f}", ""])
            total += subtotal
        filas.append(["", "TOTAL DISTRIBUIBLE", "", "", "", "", f"{total:.2f}", ""])
        tablas.append((edificio["nombre"], filas))
        for f in reporte.get("fondos_moras", []):
            if f["edificio_id"] == edificio["id"]:
                tablas.append((f"{edificio['nombre']} — fondos de moras al capturar el reporte", [
                    ["Concepto", "Importe"],
                    ["Disponible inicial del periodo", f["saldo_disponible_inicial"]],
                    ["Cobros netos del periodo", f["cobros_netos_periodo"]],
                    ["Cobros reconocidos del periodo", f["cobros_confirmados_periodo"]],
                    ["Reversos de cobros del periodo", f["reversos_cobro_periodo"]],
                    ["Coberturas confirmadas (ya reservadas)", f["aplicaciones_confirmadas_periodo"]],
                    ["Reservas del periodo", f["reservas_periodo"]],
                    ["Liberaciones del periodo", f["liberaciones_periodo"]],
                    ["Disponible final del periodo", f["saldo_disponible_final"]],
                    ["Disponible al capturar", f["saldo_disponible"]],
                    ["Reservado al capturar", f["saldo_reservado"]],
                    ["Total aplicado", f["total_aplicado"]],
                    ["Mora pendiente de cobro (no es caja)", f["mora_pendiente_cobro"]],
                ]))
        cargos = [["Unidad", "Cargo definitivo"]] + [[numero, reporte["cuotas"][id]] for id, numero in numeros.items() if id in reporte.get("cuotas", {})]
        if len(cargos) > 1:
            tablas.append((f"{edificio['nombre']} — cargos por unidad", cargos))
        if len(ajustes) > 1:
            tablas.append((f"{edificio['nombre']} — cierre de centavos", ajustes))
    return tablas


def generar_documento(reporte, formato):
    titulo = f"Prorrateo mantenimiento {reporte['periodo']} — {reporte['estado']}"
    metadata = f"Versión {reporte['version']} · Moneda {reporte['moneda']} · Vencimiento {reporte['fecha_vencimiento']}"
    nota = "Unitario de referencia redondeado. Los cargos definitivos incluyen el cierre de centavos."
    aprobacion = f"Aprobación: {reporte.get('aprobado_por') or 'Pendiente'} · {reporte.get('aprobado_en') or ''}"
    tablas = tablas_reporte(reporte)
    salida = BytesIO()
    if formato == "docx":
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Cm

        documento = Document()
        seccion = documento.sections[0]
        seccion.orientation = WD_ORIENT.LANDSCAPE
        seccion.page_width, seccion.page_height = Cm(29.7), Cm(21)
        seccion.left_margin = seccion.right_margin = Cm(1.5)
        documento.core_properties.title = titulo
        documento.add_heading(titulo, 0)
        for texto in (metadata, aprobacion, nota):
            documento.add_paragraph(texto)
        for nombre, filas in tablas:
            documento.add_heading(nombre, level=1)
            tabla = documento.add_table(rows=0, cols=len(filas[0]))
            tabla.style = "Table Grid"
            for fila in filas:
                for celda, texto in zip(tabla.add_row().cells, fila):
                    celda.text = texto
        documento.save(salida)
    elif formato == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

        estilos = getSampleStyleSheet()
        normal = estilos["BodyText"]
        normal.fontSize, normal.leading = 8, 10
        contenido = [Paragraph(escape(titulo), estilos["Title"])]
        contenido += [Paragraph(escape(texto), normal) for texto in (metadata, aprobacion, nota)]
        for nombre, filas in tablas:
            contenido += [Spacer(1, 12), Paragraph(escape(nombre), estilos["Heading2"])]
            anchos = [50, 180, 110, 75, 85, 85, 85, 75] if len(filas[0]) == 8 else None
            tabla = LongTable([[Paragraph(escape(t), normal) for t in fila] for fila in filas], repeatRows=1, colWidths=anchos)
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            contenido.append(tabla)
        SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=30, rightMargin=30,
                          topMargin=30, bottomMargin=30, title=titulo).build(contenido)
    else:
        raise ValueError("Formato de exportación no admitido.")
    return salida.getvalue()
