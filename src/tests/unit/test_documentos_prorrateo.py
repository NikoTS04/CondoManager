from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from src.modules.cuotas.documentos_prorrateo import generar_documento


def reporte():
    return {
        "periodo": "2026-11",
        "estado": "BORRADOR",
        "version": 3,
        "moneda": "PEN",
        "fecha_vencimiento": "2026-11-10",
        "edificios": [
            {"id": "e", "nombre": "Edificio 03", "departamentos": [{"id": "d", "numero": "301"}]}
        ],
        "cuotas": {"d": "500.00"},
        "items": [
            {
                "codigo": "MAN-01",
                "descripcion": "Mantenimiento <ascensores> & revisión",
                "seccion": "EDIFICIO",
                "metodo_asignacion": "DIRECTO_EDIFICIO",
                "monto_periodo": "500.00",
                "tipo_movimiento": "CARGO",
                "asignaciones": [
                    {
                        "edificio_id": "e",
                        "monto": "500.00",
                        "unitario": "500.00",
                        "cuotas": [{"departamento_id": "d", "ajuste": "0.00"}],
                    }
                ],
            }
        ],
    }


def test_ambos_formatos_conservan_importes_version_y_borrador():
    data = reporte()
    pdf = generar_documento(data, "pdf")
    docx = generar_documento(data, "docx")
    texto_pdf = " ".join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    documento = Document(BytesIO(docx))
    texto_docx = " ".join(
        [p.text for p in documento.paragraphs]
        + [c.text for t in documento.tables for r in t.rows for c in r.cells]
    )
    for texto in (texto_pdf, texto_docx):
        for esperado in ("BORRADOR", "Versión 3", "MAN-01", "500.00", "301", "<ascensores>"):
            assert esperado in texto
    assert data == reporte()


def test_pdf_multipagina_y_formato_invalido():
    data = reporte()
    data["items"] *= 150
    assert len(PdfReader(BytesIO(generar_documento(data, "pdf"))).pages) > 1
    with pytest.raises(ValueError):
        generar_documento(data, "exe")
