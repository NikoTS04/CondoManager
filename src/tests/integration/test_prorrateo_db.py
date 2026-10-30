"""Detalle, aprobación, emisión y privacidad sobre PostgreSQL migrado real."""

import json
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document
from jsonschema import Draft202012Validator
from pypdf import PdfReader
from sqlalchemy import select

from src.modules.condominios.models import Departamento, Edificio
from src.modules.cuotas.models import CuotaMantenimiento, EventoLoteOutbox, LoteCuotas
from src.modules.cuotas.presupuestos.models import ConceptoPresupuestal
from src.modules.notificaciones.consumo import consumir
from src.tests.integration.test_pagos_db import login

pytestmark = [pytest.mark.asyncio, pytest.mark.postgres]


async def presupuesto(client, base, *, cero=False):
    junta = await login(client, "admin", "Admin123!")
    creado = await client.post(
        f"/api/v1/condominios/{base.cid}/conceptos-presupuestales",
        headers=junta,
        json={
            "codigo": "MAN-01",
            "categoria": "MANTENIMIENTO_PREVENTIVO",
            "descripcion": "Ascensores",
        },
    )
    assert creado.status_code == 201, creado.text
    concepto_id = creado.json()["id"]
    item = {
        "concepto_id": concepto_id,
        "seccion": "AACC",
        "monto_periodo": "4800.00",
        "referencia_documento": "SOLO-ADMIN",
    }
    items = [item, {**item, "tipo_movimiento": "ABONO"}] if cero else [item]
    r = await client.post(
        "/api/v1/presupuestos",
        headers=junta,
        json={
            "condominio_id": base.cid,
            "periodo": "2026-11",
            "moneda": "PEN",
            "fecha_vencimiento": "2026-11-10",
            "items": items,
        },
    )
    assert r.status_code == 201, r.text
    return junta, r.json(), concepto_id


@pytest.mark.parametrize("cero", [False, True])
async def test_aprobacion_emision_y_snapshot_cero_o_positivo(api_pagos, cero):
    client, base = api_pagos
    junta, p, concepto = await presupuesto(client, base, cero=cero)
    assert p["monto_total"] == ("0.00" if cero else "4800.00")
    assert p["version"] == 1
    sinversion = await client.post(f"/api/v1/presupuestos/{p['id']}/aprobar", headers=junta)
    assert sinversion.status_code == 409
    r = await client.post(
        f"/api/v1/presupuestos/{p['id']}/aprobar?version_esperada=1", headers=junta
    )
    assert r.status_code == 200 and r.json()["version"] == 2, r.text
    async with base.sessions() as db, db.begin():
        edificios = list(await db.scalars(select(Edificio)))
        eid = edificios[0].id
        c = await db.get(ConceptoPresupuestal, UUID(concepto))
        c.descripcion = "Nombre modificado después de aprobar"
        antes = {str(d.id): d.saldo_a_favor for d in await db.scalars(select(Departamento))}
    reporte = await client.get(f"/api/v1/presupuestos/{p['id']}/prorrateo", headers=junta)
    assert reporte.status_code == 200 and reporte.json()["items"][0]["descripcion"] == "Ascensores"
    residente = await login(client)
    privado = await client.get(
        f"/api/v1/presupuestos/{p['id']}/prorrateo?edificio_id={eid}", headers=residente
    )
    assert privado.status_code == 200, privado.text
    assert "cuotas" not in privado.json()
    assert "departamentos" not in privado.json()["edificios"][0]
    assert "cuotas" not in privado.json()["items"][0]["asignaciones"][0]
    assert "SOLO-ADMIN" not in privado.text
    descarga = await client.get(
        f"/api/v1/presupuestos/{p['id']}/prorrateo/exportar?version=2&formato=docx&edificio_id={eid}",
        headers=residente,
    )
    assert descarga.status_code == 200
    documento = Document(BytesIO(descarga.content))
    texto = " ".join(c.text for t in documento.tables for r in t.rows for c in r.cells)
    assert "Cargo definitivo" not in texto and "SOLO-ADMIN" not in texto
    emitido = await client.post(
        "/api/v1/cuotas/emitir-lote",
        headers=junta,
        json={"condominio_id": base.cid, "periodo": "2026-11"},
    )
    assert emitido.status_code == 201, emitido.text
    async with base.sessions() as db:
        cuotas = list(
            await db.scalars(
                select(CuotaMantenimiento).where(CuotaMantenimiento.periodo == "2026-11")
            )
        )
        assert sum(c.monto_total_exigible for c in cuotas) == Decimal(p["monto_total"])
        assert all(c.detalle_prorrateo for c in cuotas)
        if cero:
            assert all(c.estado == "PAGADA" and c.monto_pagado == 0 for c in cuotas)
            assert {
                str(d.id): d.saldo_a_favor for d in await db.scalars(select(Departamento))
            } == antes
        lote = await db.scalar(select(LoteCuotas).where(LoteCuotas.periodo == "2026-11"))
        assert lote.metodo == "JERARQUICA"
        evento = await db.scalar(
            select(EventoLoteOutbox).where(EventoLoteOutbox.lote_id == lote.id)
        )
        payload = evento.payload
    schema = json.loads(
        Path(
            "specs/03-contracts/events/schemas/cuotas-prorrateo-emitidas.v2.schema.json"
        ).read_text()
    )
    Draft202012Validator(schema).validate(payload)
    assert await consumir(base.sessions, payload) is not None
    assert await consumir(base.sessions, payload) == 0


async def test_versionado_borrador_y_padron_modificado(api_pagos):
    client, base = api_pagos
    junta, p, _ = await presupuesto(client, base)
    body = {k: p[k] for k in ("periodo", "moneda", "fecha_vencimiento", "items")}
    actualizado = await client.put(
        f"/api/v1/presupuestos/{p['id']}", headers=junta, json={**body, "version_esperada": 1}
    )
    assert actualizado.status_code == 200 and actualizado.json()["version"] == 2
    assert (
        await client.put(
            f"/api/v1/presupuestos/{p['id']}", headers=junta, json={**body, "version_esperada": 1}
        )
    ).status_code == 409
    assert (
        await client.post(
            f"/api/v1/presupuestos/{p['id']}/aprobar?version_esperada=2", headers=junta
        )
    ).status_code == 200
    async with base.sessions() as db, db.begin():
        d = await db.scalar(select(Departamento).where(Departamento.activo.is_(True)).limit(1))
        d.activo = False
    emitido = await client.post(
        "/api/v1/cuotas/emitir-lote",
        headers=junta,
        json={"condominio_id": base.cid, "periodo": "2026-11"},
    )
    assert (
        emitido.status_code == 409 and emitido.json()["error_code"] == "PADRON_APROBADO_MODIFICADO"
    )
    async with base.sessions() as db:
        assert (
            list(
                await db.scalars(
                    select(CuotaMantenimiento.id).where(CuotaMantenimiento.periodo == "2026-11")
                )
            )
            == []
        )


async def test_catalogo_duplicado_y_residente_sin_permiso(api_pagos):
    client, base = api_pagos
    junta, p, _ = await presupuesto(client, base)
    request = {"codigo": "MAN-01", "categoria": "MANTENIMIENTO_PREVENTIVO", "descripcion": "Otro"}
    assert (
        await client.post(
            f"/api/v1/condominios/{base.cid}/conceptos-presupuestales", headers=junta, json=request
        )
    ).status_code == 409
    residente = await login(client)
    assert (
        await client.get(f"/api/v1/presupuestos/{p['id']}/prorrateo", headers=residente)
    ).status_code == 403
    assert (
        await client.post(
            f"/api/v1/condominios/{base.cid}/conceptos-presupuestales",
            headers=residente,
            json=request,
        )
    ).status_code == 403


async def test_exportar_version_privada_y_error_sin_mutar(api_pagos, monkeypatch):
    client, base = api_pagos
    junta, p, _ = await presupuesto(client, base)
    ruta = f"/api/v1/presupuestos/{p['id']}/prorrateo/exportar"
    assert (await client.get(ruta, headers=junta)).status_code == 422
    assert (await client.get(ruta + "?version=99", headers=junta)).status_code == 409
    assert (await client.get(ruta + "?version=1&formato=exe", headers=junta)).status_code == 422
    for formato in ("pdf", "docx"):
        r = await client.get(ruta + f"?version=1&formato={formato}", headers=junta)
        assert r.status_code == 200, r.text
        assert r.headers["cache-control"] == "no-store"
        assert r.headers["content-disposition"].endswith(f'.{formato}"')
        if formato == "pdf":
            assert "BORRADOR" in PdfReader(BytesIO(r.content)).pages[0].extract_text()
        else:
            assert "BORRADOR" in Document(BytesIO(r.content)).paragraphs[0].text
    residente = await login(client)
    assert (await client.get(ruta + "?version=1", headers=residente)).status_code == 403

    def fallar(*_):
        raise RuntimeError("Fallo simulado")

    monkeypatch.setattr("src.modules.cuotas.prorrateo_api.generar_documento", fallar)
    assert (await client.get(ruta + "?version=1", headers=junta)).status_code == 503
    actual = await client.get(f"/api/v1/condominios/{base.cid}/presupuestos/2026-11", headers=junta)
    assert actual.json()["version"] == 1 and actual.json()["estado"] == "BORRADOR"
